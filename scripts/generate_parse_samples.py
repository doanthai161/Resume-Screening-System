"""Generate synthetic (non-personal) PDF/DOCX acceptance documents."""

import argparse
import json
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

TEXT = [
    "SYNTHETIC CV - TEST DATA",
    "Candidate: Sample Engineer",
    "Skills: Python, FastAPI, MongoDB, Redis, Docker",
    "Experience: Backend developer from 2020 to 2024.",
    "Built recruitment APIs and automated document processing.",
    "Education: Bachelor of Computer Science, Example University.",
    "Languages: Vietnamese and English.",
]


def digital_pdf(lines=TEXT) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    ),
                }
            )
        }
    )
    stream = DecodedStreamObject()
    escaped = [
        line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        for line in lines
    ]
    stream.set_data(
        (
            "BT /F1 11 Tf 40 750 Td "
            + " 0 -22 Td ".join(f"({line}) Tj" for line in escaped)
            + " ET"
        ).encode("ascii")
    )
    # PDF streams must be indirect objects. Direct assignment is tolerated by
    # pypdf text extraction but renders a blank page in PDFium/MinerU.
    page.replace_contents(stream)
    data = BytesIO()
    writer.write(data)
    return data.getvalue()


def scanned_pdf(*, vietnamese=False) -> bytes:
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=24)
    lines = TEXT + ["Scan verification marker: SCANPAGE"]
    if vietnamese:
        fonts = [
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        ]
        font_path = next((path for path in fonts if path.is_file()), None)
        if font_path is None:
            raise RuntimeError(
                "Install Arial or DejaVu Sans for the Vietnamese scan fixture"
            )
        font = ImageFont.truetype(str(font_path), size=24)
        lines += [
            "Kỹ năng: Python và MongoDB.",
            "Học vấn: Công nghệ thông tin.",
            "Kinh nghiệm phát triển phần mềm.",
        ]
    draw.multiline_text((65, 100), "\n\n".join(lines), font=font, fill="black")
    data = BytesIO()
    image.save(data, format="PDF", resolution=150)
    return data.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Hồ sơ mẫu - dữ liệu giả lập"
    for line in TEXT:
        document.add_paragraph(line)
    document.add_paragraph(
        "Kỹ năng: Python. Kinh nghiệm phát triển hệ thống tuyển dụng."
    )
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Học vấn"
    table.cell(0, 1).text = "Công nghệ thông tin"
    document.sections[0].footer.paragraphs[0].text = "FOOTER_MARKER"
    data = BytesIO()
    document.save(data)
    return data.getvalue()


def generate(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    native, scan = digital_pdf(), scanned_pdf()
    writer = PdfWriter()
    for data in (native, scan):
        writer.append(PdfReader(BytesIO(data)))
    mixed = BytesIO()
    writer.write(mixed)
    files = {
        "digital.pdf": native,
        "scan.pdf": scan,
        "scan_vi.pdf": scanned_pdf(vietnamese=True),
        "mixed.pdf": mixed.getvalue(),
        "vietnamese.docx": docx_bytes(),
    }
    cases = []
    for name, data in files.items():
        with (output / name).open("xb") as handle:
            handle.write(data)
        terms = ["Python", "MongoDB", "Computer Science"]
        if name in {"scan.pdf", "scan_vi.pdf", "mixed.pdf"}:
            terms.append("SCANPAGE")
        if name == "scan_vi.pdf":
            terms.extend(["Kỹ năng", "Học vấn"])
        if name.endswith("docx"):
            terms.extend(["Kỹ năng", "Học vấn", "FOOTER_MARKER"])
        cases.append(
            {
                "file": name,
                "expected_terms": terms,
                "expected_ocr": name in {"scan.pdf", "scan_vi.pdf", "mixed.pdf"},
            }
        )
    with (output / "manifest.json").open("x", encoding="utf-8") as handle:
        json.dump(cases, handle, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    generate(parser.parse_args().output)
