import asyncio
import hashlib
import json
import sys
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from docx.oxml.ns import qn

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import settings
from app.models.resume_file import ResumeFile
from app.models.screening_run import ResumeParseRun
from app.workers.errors import PermanentParseError, UnsafeDocumentError


SUPPORTED_MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
MAX_EXTRACTED_CHARACTERS = 1_000_000


@dataclass(frozen=True)
class ResumeDocument:
    resume: ResumeFile
    content: bytes
    extension: str
    filename: str


@dataclass(frozen=True)
class DocumentAnalysis:
    page_count: int
    native_text: str
    text_page_ratio: float
    requires_ocr: bool
    native_complete: bool = True
    header_text: str = ""
    footer_text: str = ""


def pdf_ocr_mode(analysis: DocumentAnalysis) -> str:
    needs_ocr = analysis.requires_ocr or analysis.text_page_ratio < 1.0
    if needs_ocr and not settings.MINERU_AUTO_OCR:
        raise PermanentParseError("This PDF requires OCR", code="pdf_ocr_required")
    return "ocr" if needs_ocr else "txt"


def _read_local_resume(resume: ResumeFile) -> ResumeDocument:
    if resume.storage_provider != "local":
        raise PermanentParseError(
            "Only local resume storage is supported",
            code="unsupported_storage_provider",
        )
    try:
        root = settings.resume_upload_path.resolve()
        path = Path(resume.file_path).resolve(strict=True)
    except (OSError, RuntimeError) as exc:
        raise UnsafeDocumentError("Resume file cannot be resolved") from exc
    extension = path.suffix.lower()
    if not path.is_relative_to(root) or not path.is_file() or extension not in SUPPORTED_MIME_TYPES:
        raise UnsafeDocumentError("Resume path or format is invalid")
    if resume.mime_type != SUPPORTED_MIME_TYPES[extension]:
        raise UnsafeDocumentError("Resume MIME type does not match its extension")
    if not 0 < resume.file_size <= settings.MAX_RESUME_SIZE:
        raise UnsafeDocumentError("Resume size metadata is invalid")
    try:
        with path.open("rb") as handle:
            content = handle.read(settings.MAX_RESUME_SIZE + 1)
    except OSError as exc:
        raise UnsafeDocumentError("Resume file cannot be read") from exc
    if len(content) != resume.file_size:
        raise UnsafeDocumentError("Resume size does not match stored metadata")
    if hashlib.sha256(content).hexdigest() != resume.checksum:
        raise UnsafeDocumentError("Resume checksum does not match stored metadata")
    return ResumeDocument(
        resume=resume,
        content=content,
        extension=extension,
        filename=path.name,
    )


async def load_resume_document(run: ResumeParseRun) -> ResumeDocument:
    resume = await ResumeFile.find_one(
        {
            "_id": run.resume_file_id,
            "company_id": run.company_id,
            "is_deleted": False,
        }
    )
    if resume is None:
        raise PermanentParseError(
            "Resume is missing or belongs to another company",
            code="resume_not_found",
        )
    return await asyncio.to_thread(_read_local_resume, resume)


def _join_text(parts: list[str]) -> str:
    text = "\n\n".join(part.strip() for part in parts if part.strip())
    # Never mark a truncated CV complete: downstream extraction cannot detect
    # that the end of the applicant's experience was silently removed.
    if len(text) > MAX_EXTRACTED_CHARACTERS or len(json.dumps(text).encode("utf-8")) > 1_900_000:
        raise PermanentParseError("Resume text exceeds the limit", code="document_text_too_large")
    return text


def _analyze_pdf(document: ResumeDocument) -> DocumentAnalysis:
    try:
        reader = PdfReader(BytesIO(document.content), strict=False)
        if reader.is_encrypted:
            raise PermanentParseError(
                "Password-protected PDFs are not supported",
                code="encrypted_pdf",
            )
        page_count = len(reader.pages)
        if not 0 < page_count <= settings.MAX_RESUME_PAGES:
            raise PermanentParseError(
                "PDF page count is outside the allowed range",
                code="invalid_page_count",
            )
    except PermanentParseError:
        raise
    except Exception as exc:
        # Unknown page count cannot safely bypass MAX_RESUME_PAGES.
        raise PermanentParseError("PDF structure cannot be inspected", code="invalid_pdf") from exc

    page_texts = []
    total = 0
    for page in reader.pages:
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""  # Known page count, but this page requires OCR.
        total += len(text) + 2
        if total > MAX_EXTRACTED_CHARACTERS:
            raise PermanentParseError("Resume text exceeds the limit", code="document_text_too_large")
        page_texts.append(text)

    useful_pages = sum(
        len(text) >= settings.PARSE_MIN_TEXT_CHARACTERS_PER_PAGE
        for text in page_texts
    )
    ratio = useful_pages / page_count
    return DocumentAnalysis(
        page_count=page_count,
        native_text=_join_text(page_texts),
        text_page_ratio=round(ratio, 4),
        requires_ocr=useful_pages < page_count,
    )


def _analyze_docx(document: ResumeDocument) -> DocumentAnalysis:
    try:
        # Recheck archive limits for old/imported records, before XML expansion.
        from app.services.resume_service import _validate_docx_archive
        if not _validate_docx_archive(BytesIO(document.content)):
            raise UnsafeDocumentError("DOCX archive is invalid or exceeds limits")
        source = DocxDocument(BytesIO(document.content))
        roots = [("body", source.element.body)]
        for part in source.part.related_parts.values():
            if part.content_type.endswith("wordprocessingml.header+xml"):
                roots.append(("header", part.element))
            elif part.content_type.endswith("wordprocessingml.footer+xml"):
                roots.append(("footer", part.element))
        parts = []
        margins = {"header": [], "footer": []}
        native_complete = True
        for kind, root in roots:
            # XML traversal preserves paragraph/table/textbox order, including
            # nested tables. It also reads headers/footers omitted by .paragraphs.
            root_parts = []
            for element in root.iter():
                if element.tag == qn("w:t"):
                    root_parts.append(element.text or "")
                elif element.tag in {qn("w:p"), qn("w:br"), qn("w:tab")}:
                    root_parts.append("\n")
                elif element.tag in {qn("w:drawing"), qn("w:pict"), qn("w:altChunk"), qn("w:object")}:
                    native_complete = False
            parts.extend(root_parts)
            if kind in margins:
                margins[kind].append("".join(root_parts))
        text = _join_text(["".join(parts)])
        header_text = _join_text(margins["header"])
        footer_text = _join_text(margins["footer"])
    except PermanentParseError:
        raise
    except Exception as exc:
        raise PermanentParseError("DOCX structure cannot be inspected", code="invalid_docx") from exc
    return DocumentAnalysis(
        page_count=1,
        native_text=text,
        text_page_ratio=1.0 if text else 0.0,
        requires_ocr=False,
        native_complete=native_complete,
        header_text=header_text,
        footer_text=footer_text,
    )


async def analyze_document(document: ResumeDocument) -> DocumentAnalysis:
    # Native libraries can keep running after to_thread is cancelled. A child
    # process gives worker timeout/lease loss a real stop boundary.
    limits = {
        "MAX_RESUME_SIZE": settings.MAX_RESUME_SIZE,
        "MAX_RESUME_PAGES": settings.MAX_RESUME_PAGES,
        "PARSE_MIN_TEXT_CHARACTERS_PER_PAGE": settings.PARSE_MIN_TEXT_CHARACTERS_PER_PAGE,
        "MAX_DOCX_ENTRIES": settings.MAX_DOCX_ENTRIES,
        "MAX_DOCX_UNCOMPRESSED_SIZE": settings.MAX_DOCX_UNCOMPRESSED_SIZE,
    }
    process = await asyncio.create_subprocess_exec(
        sys.executable, "-m", "app.workers.document_probe", document.extension,
        json.dumps(limits), stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.DEVNULL,
        cwd=str(Path(__file__).resolve().parents[2]),
    )
    try:
        try:
            data, _ = await asyncio.wait_for(
                process.communicate(document.content), settings.PARSE_PREFLIGHT_TIMEOUT_SECONDS,
            )
        except TimeoutError:
            raise PermanentParseError("Document inspection timed out", code="document_inspection_timeout") from None
        if process.returncode != 0:
            raise PermanentParseError("Document inspection failed", code="document_inspection_failed")
        result = json.loads(data)
        if "error_code" in result:
            raise PermanentParseError("Document inspection rejected the file", code=result["error_code"])
        return DocumentAnalysis(**result)
    finally:
        if process.returncode is None:
            try:
                process.kill()
            except ProcessLookupError:
                pass
        # Drain pipes as well as reaping: wait() alone can hang on a full pipe.
        await process.communicate()
